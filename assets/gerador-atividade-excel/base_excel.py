"""Gera a planilha de base (.xlsx) que o aluno baixa para começar a atividade."""

import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from desenho_excel import PADRAO_REFERENCIA, valor_numerico

FORMATO_DATA = "dd/mm/yyyy"
FORMATO_MOEDA = '"R$" #,##0.00'
LARGURA_COLUNA_PADRAO = 16


def converter_valor(texto):
    """Converte o texto da planilha do JSON no tipo certo para o Excel.

    Args:
        texto: Valor como aparece no JSON (ex.: "05/01/2026", "18,90", "P-001").

    Returns:
        datetime, float, int ou o próprio texto.
    """
    texto = str(texto)
    try:
        return datetime.strptime(texto, "%d/%m/%Y")
    except ValueError:
        pass
    numero = valor_numerico(texto)
    if numero is None or texto.startswith("P-") or PADRAO_REFERENCIA.match(texto):
        return texto
    return int(numero) if numero.is_integer() and "," not in texto else numero


def formatar_linha(aba, linha, colunas_moeda):
    """Aplica formato de data e de moeda às células de uma linha de dados.

    Args:
        aba: Planilha do openpyxl.
        linha: Número da linha.
        colunas_moeda: Letras das colunas em moeda.
    """
    for celula in aba[linha]:
        if isinstance(celula.value, datetime):
            celula.number_format = FORMATO_DATA
        elif celula.column_letter in colunas_moeda:
            celula.number_format = FORMATO_MOEDA


def gerar_base_xlsx(atividade, especificacao, pasta):
    """Grava o .xlsx de base, com título mesclado e linha vazia opcionais (para corrigir).

    Args:
        atividade: Atividade completa (planilhas do JSON).
        especificacao: {"nome", "aba", "planilha", "colunas", "titulo", "linha_vazia_apos",
            "colunas_moeda"}.
        pasta: Pasta da atividade.

    Returns:
        Caminho do arquivo gerado.
    """
    linhas = atividade["planilhas"][especificacao["planilha"]]["linhas"]
    indices = [ord(coluna) - 65 for coluna in especificacao["colunas"]]
    livro = Workbook()
    aba = livro.active
    aba.title = especificacao.get("aba", "Planilha1")
    if especificacao.get("titulo"):
        aba.append([especificacao["titulo"]])
        ultima = chr(64 + len(indices))
        aba.merge_cells(f"A1:{ultima}1")
        aba["A1"].font = Font(bold=True, size=13)
        aba["A1"].alignment = Alignment(horizontal="center")
    aba.append([linhas[0][i] for i in indices])
    for numero, linha in enumerate(linhas[1:], start=1):
        aba.append([converter_valor(linha[i]) for i in indices])
        formatar_linha(aba, aba.max_row, especificacao.get("colunas_moeda", []))
        if numero == especificacao.get("linha_vazia_apos"):
            aba.append([])
    for indice in range(len(indices)):
        aba.column_dimensions[chr(65 + indice)].width = LARGURA_COLUNA_PADRAO
    caminho = os.path.join(pasta, especificacao["nome"])
    livro.save(caminho)
    return caminho
